"""Text out of uploaded documents, for the AI features that read them (the
scope of work, the PO). Extracted in the AI worker when first needed and kept
in doc_text, so a document is read once. Uploads are untrusted: PDFs are read
to 40 pages, Office files to 20 MB unpacked, and the text is capped at
200 000 characters. Each file is read in a child process that is stopped
after a minute, so one pathological file cannot stall the scheduler's
housekeeping or the AI worker. Scanned PDFs (pictures of pages) have no text.

PDF: pypdf. Excel: openpyxl. Word: the document's XML, read with the
standard library. Text, CSV and email (.eml): decoded as text."""
import os
import re
import zipfile
from datetime import datetime
from xml.etree import ElementTree

MAX_CHARS, MAX_PAGES, MAX_UNPACKED = 200_000, 40, 20 * 1024 * 1024
# ponytail: fixed guards against hostile files, not business rules; no memory cap (RLIMIT_AS is unreliable off Linux)
READ_SECONDS, MAX_FILE = 60, 100 * 1024 * 1024   # the upload ceiling is 100 MB
READABLE = ("pdf", "docx", "xlsx", "txt", "csv", "eml")


def _pdf(path):
    import logging
    from pypdf import PdfReader
    logging.getLogger("pypdf").setLevel(logging.ERROR)   # a damaged upload is a document without text, not a log line per page
    reader = PdfReader(path)
    return "\n\n".join((page.extract_text() or "") for page in reader.pages[:MAX_PAGES])


def _zip_member(path, name):
    with zipfile.ZipFile(path) as z:
        info = z.getinfo(name)
        if info.file_size > MAX_UNPACKED:
            raise ValueError("the file unpacks to more than 20 MB")
        return z.read(name)


def _docx(path):
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    root = ElementTree.fromstring(_zip_member(path, "word/document.xml"))
    lines = []
    for p in root.iter(f"{w}p"):   # paragraphs, including those inside tables
        text = "".join(t.text or "" for t in p.iter(f"{w}t"))
        if text.strip():
            lines.append(text)
    return "\n".join(lines)


def _xlsx(path):
    from openpyxl import load_workbook
    with zipfile.ZipFile(path) as z:
        if sum(i.file_size for i in z.infolist()) > MAX_UNPACKED:
            raise ValueError("the file unpacks to more than 20 MB")
    wb = load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"[{ws.title}]")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c not in (None, "")]
            if cells:
                out.append(" | ".join(cells))
            if sum(len(x) for x in out) > MAX_CHARS:
                break
    wb.close()
    return "\n".join(out)


def _eml(path):
    from email import policy
    from email.parser import BytesParser
    with open(path, "rb") as fh:
        msg = BytesParser(policy=policy.default).parse(fh)
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body else ""
    if body and body.get_content_type() == "text/html":
        text = re.sub(r"<[^>]+>", " ", text)
    return f"From: {msg['from']}\nSubject: {msg['subject']}\n\n{text}"


def extract_file(path, ext):
    """(text, error): the document's text, or why there is none."""
    ext = ext.lower().lstrip(".")
    if ext not in READABLE:
        return "", f".{ext} files are not read"
    try:
        if ext == "pdf":
            text = _pdf(path)
        elif ext == "docx":
            text = _docx(path)
        elif ext == "xlsx":
            text = _xlsx(path)
        elif ext == "eml":
            text = _eml(path)
        else:
            with open(path, "rb") as fh:
                text = fh.read(MAX_CHARS * 4).decode("utf-8", errors="replace")
    except Exception as exc:  # noqa: BLE001 — a broken or hostile file is a document without text, never a crash
        return "", f"could not be read ({type(exc).__name__})"
    text = re.sub(r"[ \t]+", " ", re.sub(r"\n{3,}", "\n\n", text)).strip()[:MAX_CHARS]
    if not text:
        return "", "no text in it (a scan or a picture?)"
    return text, None


def extract_isolated(path, ext):
    """extract_file in a child process with a time limit (review P3-29): a
    file that would loop or eat the CPU is stopped, not waited for."""
    import json
    import subprocess
    import sys
    if os.path.getsize(path) > MAX_FILE:
        return "", "too large to read"
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    try:
        done = subprocess.run([sys.executable, "-m", "services.doctext", path, ext], cwd=root, capture_output=True, timeout=READ_SECONDS)
        text, error = json.loads(done.stdout)
    except subprocess.TimeoutExpired:
        return "", f"took more than {READ_SECONDS} seconds to read"
    except ValueError:   # the child died (out of memory, a crash) before it answered
        return "", "could not be read (the reader stopped)"
    return text, error


def extract_pending(db=None, limit=200):
    """Read the text of documents not read yet, so search finds words in them.
    Nightly (scheduler housekeeping), a few hundred at a time."""
    from db import get_db
    db = db or get_db()
    ids = [r[0] for r in db.execute("SELECT d.id FROM documents d LEFT JOIN doc_text t ON t.document_id = d.id "
                                    "WHERE d.deleted_at IS NULL AND d.review_status = 'accepted' AND t.document_id IS NULL ORDER BY d.id DESC LIMIT ?", (limit,))]
    for i in ids:
        text_of(i, db)
    return len(ids)


def text_of(document_id, db=None):
    """A document's text, extracted on first use and kept. (text, error)."""
    from db import UPLOAD_DIR, get_db
    db = db or get_db()
    row = db.execute("SELECT text, error FROM doc_text WHERE document_id = ?", (document_id,)).fetchone()
    if row:
        return row["text"] or "", row["error"]
    doc = db.execute("SELECT id, project_id, filename, stored_name, review_status FROM documents WHERE id = ? AND deleted_at IS NULL", (document_id,)).fetchone()
    if doc is None:
        return "", "the document is gone"
    if doc["review_status"] != "accepted":   # an outside upload is not read before staff accept it (services/doc_review.py)
        return "", "not reviewed yet"
    path = os.path.join(UPLOAD_DIR, str(doc["project_id"]), doc["stored_name"])
    text, error = extract_isolated(path, os.path.splitext(doc["filename"])[1]) if os.path.exists(path) else ("", "the file is missing")
    db.execute("INSERT OR REPLACE INTO doc_text (document_id, project_id, chars, text, error, extracted_at) VALUES (?, ?, ?, ?, ?, ?)",
               (doc["id"], doc["project_id"], len(text), text, error, datetime.utcnow().isoformat()))
    db.commit()
    return text, error


if __name__ == "__main__":   # the child of extract_isolated: python -m services.doctext <path> <ext>
    import json
    import sys
    sys.stdout.write(json.dumps(extract_file(sys.argv[1], sys.argv[2])))
